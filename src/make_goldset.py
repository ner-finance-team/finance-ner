# -*- coding: utf-8 -*-
"""
2단계 금융 골드셋 라벨링 파일 만들기 (라벨 v1, 팀 결정 전 제안안 기준).
  1) 경제 문장 전체(data/econ_pool/econ_all.jsonl)에서 20~200글자 문장 20만 개를 무작위로 뽑는다.
  2) 선라벨: MONEY·RATE·DATE는 규칙, INDICATOR는 지표 사전 v0, FIN_ORG는 이름 끝말+당국 목록, FIN_PRODUCT·POLICY는 키워드,
     ORG는 원래 데이터의 OG 중 금융기관·지수가 아닌 것. (사람이 고치기 위한 초안이다. 정답이 아니다.)
  3) 층화 추출로 510문장: 드문 라벨(POLICY·FIN_PRODUCT·FIN_ORG)이 든 문장을 더 뽑고, 선라벨이 없는 문장을 약 20% 넣는다.
  4) 배정: 맞추기 10(5명 모두) + 공통 40(5명 모두) + 개인 460(92개씩 5명).
  5) 사람별 엑셀(시트: 0_지침, 1_맞추기, 2_공통, 3_개인)과 원본 jsonl을 data/goldset/ 에 저장.
사용: python src/make_goldset.py
"""
import collections, json, os, random, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = "data/goldset"
MEMBERS = [("sejin_choi", "최세진"), ("kim_byeongseok", "김병석"), ("kim_hyeonho", "김현호"), ("jin_taehyeon", "진태현"), ("park_gyutae", "박규태")]
LABELS = ["FIN_ORG", "ORG", "FIN_PRODUCT", "MONEY", "RATE", "DATE", "INDICATOR", "POLICY"]
SEED = 1022

# ---------- 선라벨 규칙 ----------
N = r"\d[\d,]*(?:\.\d+)?"
RX = {
    "MONEY": re.compile(rf"{N}(?:\s?(?:조|억|만|천)(?:\s?{N})?)*\s?(?:원|달러|엔|위안|유로|파운드)"),
    "RATE": re.compile(rf"{N}\s?(?:%p|％p|%|％|bp|퍼센트포인트|퍼센트)"),
    "DATE": re.compile(r"(?:\d{4}년(?:\s?\d{1,2}월(?:\s?\d{1,2}일)?)?|\d{1,2}월\s?\d{1,2}일|\d{1,2}월|\d분기|(?:상|하)반기|올해|지난해|작년|내년|전년|올\s(?=\d분기|상반기|하반기))"),
}
FIN_SUFFIX = re.compile(r"[가-힣a-zA-Z0-9]{1,12}(?:금융지주|금융투자|투자증권|자산운용|저축은행|은행|증권|생명|손해보험|화재|캐피탈|카드|보험)(?![가-힣])")
AUTHORITY = ["금융위원회", "금융위", "금융감독원", "금감원", "한국은행", "한은", "예금보험공사", "예보", "한국거래소", "금융통화위원회", "금통위", "증권선물위원회", "증선위"]
PRODUCT = ["주택담보대출", "전세자금대출", "전세대출", "신용대출", "마이너스통장", "정기예금", "정기적금", "적금", "예금", "펀드", "ETF", "etf", "ETN", "etn",
           "회사채", "국채", "채권", "리츠", "개인연금", "퇴직연금", "연금저축", "연금보험", "IRP", "irp", "ISA", "isa", "CMA", "cma", "주식",
           "신용카드", "체크카드", "자동차보험", "실손보험", "실손의료보험", "종신보험", "암보험", "청약통장", "주택청약종합저축", "보금자리론", "주가연계증권", "ELS", "els", "DLS", "dls"]
NOT_FIN_ORG = {"신용카드", "체크카드", "자동차보험", "실손보험", "종신보험", "암보험", "연금보험", "실손의료보험", "법인카드", "교통카드", "선불카드", "보험", "카드", "은행", "증권"}
POLICY = ["DSR", "dsr", "LTV", "ltv", "DTI", "dti", "종합부동산세", "종부세", "양도소득세", "양도세", "재산세", "취득세", "증권거래세", "금융투자소득세", "금투세", "공매도 금지"]
POLICY_RX = re.compile(r"(?:(?:[가-힣0-9·]{2,12}\s)?[가-힣]{2,10}대책|[가-힣]{3,12}법)(?![가-힣])")
JOSA_END = re.compile(r"(?:은|는|이|가|을|를|의|에|에서|으로|로|와|과|도|만)$")
NOT_LAW = {"방법", "불법", "편법", "위법", "합법", "적법", "수법", "기법", "문법", "해법", "용법", "탈법"}
INDEX_AS_IND = {"코스피", "코스닥", "나스닥"}

def load_ind():
    d = json.load(open("data/fin_labels/indicator_dict_v0.json", encoding="utf-8"))
    return sorted({x["term"] for g in d["groups"].values() for x in g}, key=len, reverse=True)
IND = sorted(set(load_ind()) | {"영업익", "순익", "영업손익"}, key=len, reverse=True)

def josa_ok(t, e):   # 단어 안쪽 일부가 아닌지: 뒤 글자가 한글이면 흔한 조사만 허용
    nxt = t[e:e + 1]
    return not nxt or not re.match(r"[가-힣]", nxt) or re.match(r"[은는이가을를의에과와도만로며]", nxt)

def prelabel(text, og_spans):
    cands = []   # (start, end, tag, priority)
    for tag in ("MONEY", "RATE", "DATE"):
        for m in RX[tag].finditer(text): cands.append((m.start(), m.end(), tag, 3))
    for w in IND:
        for m in re.finditer(re.escape(w), text):
            if josa_ok(text, m.end()): cands.append((m.start(), m.end(), "INDICATOR", 2))
    for w in sorted(PRODUCT, key=len, reverse=True):
        for m in re.finditer(re.escape(w), text):
            if josa_ok(text, m.end()): cands.append((m.start(), m.end(), "FIN_PRODUCT", 2))
    for w in AUTHORITY:
        for m in re.finditer(re.escape(w), text):
            if josa_ok(text, m.end()): cands.append((m.start(), m.end(), "FIN_ORG", 2))
    for m in FIN_SUFFIX.finditer(text):
        if m.group(0) not in NOT_FIN_ORG and not any(m.group(0).endswith(p) for p in ("신용카드", "체크카드")):
            cands.append((m.start(), m.end(), "FIN_ORG", 1))
    for w in POLICY:
        for m in re.finditer(re.escape(w), text):
            if josa_ok(text, m.end()): cands.append((m.start(), m.end(), "POLICY", 2))
    for m in POLICY_RX.finditer(text):
        last = m.group(0).split()[-1]
        first = m.group(0).split()[0]
        if " " in m.group(0) and JOSA_END.search(first): continue      # '행위가 OO법'처럼 앞말이 조사로 끝나면 제외
        if last not in NOT_LAW and not any(last.endswith(x) for x in NOT_LAW): cands.append((m.start(), m.end(), "POLICY", 1))
    for s, e, t in og_spans:
        if t == "OG":
            w = text[s:e]
            fin = w in AUTHORITY or (FIN_SUFFIX.fullmatch(w) and w not in NOT_FIN_ORG)
            cands.append((s, e, "INDICATOR" if w in INDEX_AS_IND else ("FIN_ORG" if fin else "ORG"), 0))
    # 겹치면: 긴 것 → 우선순위 높은 것
    keep = []
    for s, e, t, p in sorted(cands, key=lambda x: (-(x[1] - x[0]), -x[3], x[0])):
        if all(e <= a or b <= s for a, b, _ in keep): keep.append((s, e, t))
    return sorted(keep)

def main():
    rng = random.Random(SEED); os.makedirs(OUT, exist_ok=True)
    # 1) 20~200글자 문장 20만 개 무작위(저수지 추출)
    K, pool, n = 200000, [], 0
    with open("data/econ_pool/econ_all.jsonl", encoding="utf-8") as f:
        for line in f:
            r = json.loads(line)
            if not 20 <= len(r["text"]) <= 200 or "(이름)" in r["text"][:3]: continue
            n += 1
            if len(pool) < K: pool.append(r)
            else:
                j = rng.randrange(n)
                if j < K: pool[j] = r
    print(f"후보 문장 {n:,} 중 {len(pool):,} 추출")
    for r in pool: r["prelabel"] = prelabel(r["text"], r["spans_7tag"])
    has = lambda r, t: any(x[2] == t for x in r["prelabel"])
    # 2) 층화 추출 (문장 단위 할당량)
    quota = [("POLICY", 70), ("FIN_PRODUCT", 85), ("FIN_ORG", 75), ("INDICATOR", 50), ("ORG", 30), ("MONEY", 25), ("RATE", 25), ("DATE", 20)]
    rng.shuffle(pool); used, picked = set(), []
    for tag, q in quota:
        got = 0
        for r in pool:
            if got >= q: break
            if r["id"] in used or not has(r, tag): continue
            used.add(r["id"]); picked.append(r); got += 1
    none = [r for r in pool if r["id"] not in used and not r["prelabel"]][:100]
    for r in none: used.add(r["id"])
    rest = [r for r in pool if r["id"] not in used][:510 - len(picked) - len(none)]
    sel = picked + none + rest
    assert len(sel) == 510, len(sel)
    # 3) 배정: 맞추기 10은 라벨이 다양한 문장으로
    rng.shuffle(sel)
    rich = sorted(sel, key=lambda r: -len({x[2] for x in r["prelabel"]}))[:30]
    calib = rng.sample(rich, 10); cid = {r["id"] for r in calib}
    others = [r for r in sel if r["id"] not in cid]
    shared, indiv = others[:40], others[40:]
    for i, r in enumerate(calib): r["set"], r["no"] = "1_맞추기", i + 1
    for i, r in enumerate(shared): r["set"], r["no"] = "2_공통", i + 1
    for k, (en, ko) in enumerate(MEMBERS):
        for i, r in enumerate(indiv[k * 92:(k + 1) * 92]): r["set"], r["no"], r["assignee"] = "3_개인", i + 1, en
    with open(f"{OUT}/goldset_master.jsonl", "w", encoding="utf-8") as f:
        for r in calib + shared + indiv:
            f.write(json.dumps({k: r.get(k) for k in ("id", "set", "no", "assignee", "site", "text", "prelabel", "spans_7tag")}, ensure_ascii=False) + "\n")
    # 통계: 골드셋 500(공통+개인)의 선라벨 기준 라벨별 문장 수
    gold = shared + indiv
    cov = {t: sum(has(r, t) for r in gold) for t in LABELS}
    cov["선라벨 없음"] = sum(1 for r in gold if not r["prelabel"])
    json.dump({"pool": len(pool), "candidates": n, "goldset": len(gold), "calibration": 10, "shared": 40, "individual_each": 92,
               "sentences_with_label_prelabel": cov}, open(f"{OUT}/stats.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("골드셋 500 선라벨 기준 라벨별 문장 수:", cov)
    # 4) 엑셀
    guide = [
        ["2단계 금융 골드셋 라벨링 지침 (라벨 v1 제안안 — 노션 '금융 NER 라벨 체계 초안 (정리본)' 3·4장)"],
        [""],
        ["하는 일", "각 문장의 라벨 칸에 미리 채운 선라벨(기계가 붙인 초안)을 보고 고친다. 맞으면 그대로, 틀리면 지우고, 빠졌으면 추가한다."],
        ["적는 법", "개체 글자를 문장에 나온 그대로 적고, 여러 개면 ' | '(세로 막대)로 구분. 예) 우리은행 | 신한은행"],
        ["다 봤으면", "'확인' 칸에 o 를 적는다. 애매하면 '메모' 칸에 이유를 적는다(맞추기 회의에서 같이 정함)."],
        [""],
        ["라벨", "정의", "포함", "제외", "예"],
        ["FIN_ORG", "금융회사와 금융 감독·정책 기관", "은행·증권·보험·카드·캐피탈·저축은행·자산운용·금융지주, 금융위·금감원·한국은행·예보·한국거래소", "금융업이 아닌 기업, 지점명만 단독", "우리은행, 한국투자증권, 금감원"],
        ["ORG", "그 밖의 기업·기관·단체", "일반 기업, 정부 부처, 공기업, 국회", "FIN_ORG에 속하는 것", "삼성전자, 국토교통부"],
        ["FIN_PRODUCT", "금융상품과 그 유형", "주식·펀드·ETF·채권·리츠·대출 유형·예적금·연금·카드·보험 유형, 상품명", "'대출'·'보험' 같은 일반 단어 단독(미정 — 메모)", "ETF, 회사채, 주택담보대출, 체크카드"],
        ["MONEY", "금액", "숫자+단위+통화", "앞의 약·최대·총", "5000억원, 1조원"],
        ["RATE", "비율·금리 값", "숫자+%·%p·bp", "지표 이름", "3.5%, 0.25%p"],
        ["DATE", "날짜·기간", "2026년 10월, 지난해, 올해, 3분기, 상반기", "시각", "지난해, 3분기"],
        ["INDICATOR", "경제·금융 지표, 기업 재무지표", "주가, 코스피, 코스닥, 기준금리, 환율, 성장률, 매출, 영업이익", "지표 값", "코스피, 영업이익"],
        ["POLICY", "금융·부동산 정책·규제·세제", "DSR, LTV, 양도세, OO대책, 법령명", "'규제' 단독", "DSR, 자본시장법"],
        [""],
        ["경계 규칙", "조사·어미는 뺀다(우리은행은 → 우리은행). 하나의 대상이면 여러 어절도 하나로(최장 단위). 숫자 라벨은 숫자와 단위까지만. 지표와 값은 따로(기준금리 3.5% → 기준금리 / 3.5%). 시점·지역 수식어는 빼고(지난해 영업이익 → 지난해(DATE) + 영업이익(INDICATOR))."],
        ["모르겠으면", "가장 그럴듯하게 붙이고 메모에 적는다. 비워 두지 말 것(빈칸 = 개체 없음으로 읽힘)."],
    ]
    head = ["번호", "id", "문장"] + LABELS + ["확인", "메모"]
    fill = PatternFill("solid", fgColor="FFF2CC")
    def sheet(wb, name, rows):
        ws = wb.create_sheet(name); ws.append(head)
        for c in ws[1]: c.font = Font(bold=True); c.fill = fill
        for r in rows:
            cells = {t: " | ".join(r["text"][s:e] for s, e, tt in r["prelabel"] if tt == t) for t in LABELS}
            ws.append([r["no"], r["id"], r["text"]] + [cells[t] for t in LABELS] + ["", ""])
        widths = {"A": 6, "B": 14, "C": 70}
        for col, w in widths.items(): ws.column_dimensions[col].width = w
        for i in range(len(LABELS) + 2): ws.column_dimensions[chr(ord("D") + i)].width = 16
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "D2"
    for k, (en, ko) in enumerate(MEMBERS):
        wb = Workbook(); ws = wb.active; ws.title = "0_지침"
        for row in guide: ws.append(row)
        ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 40
        for col in "CDE": ws.column_dimensions[col].width = 40
        for row in ws.iter_rows():
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        sheet(wb, "1_맞추기", calib); sheet(wb, "2_공통", shared)
        sheet(wb, "3_개인", [r for r in indiv if r.get("assignee") == en])
        wb.save(f"{OUT}/gold_{k+1}_{en}.xlsx")
    print("저장:", OUT)

if __name__ == "__main__":
    main()
