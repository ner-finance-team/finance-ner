# -*- coding: utf-8 -*-
"""
라벨링이 끝난 골드셋 엑셀(data/goldset/gold_*_이름.xlsx)을 읽어
  1) 칸에 적은 개체 글자를 문장 안 위치(span)로 바꾸고(같은 글자가 여러 번 나오면 모두, 겹치면 긴 것 우선),
  2) 맞추기·공통 시트로 사람끼리 일치율(두 사람씩 짝지은 span F1의 평균, 라벨별)을 재고,
  3) 공통 40문장은 5명 중 3명 이상이 같은 span을 적은 것을 정답으로(다수결), 개인 460문장은 그 사람 것을 정답으로 합쳐
     data/goldset/gold_final.jsonl (문장 + [시작, 끝, 태그]) 을 만든다. 채점기(evaluate.py)의 --gold 로 바로 쓴다.
문장에서 찾지 못한 글자(오타 등)는 목록으로 알려 준다.
사용: python src/import_goldset.py [엑셀 폴더(기본 data/goldset)]
"""
import collections, glob, itertools, json, os, sys
from openpyxl import load_workbook

DIR = sys.argv[1] if len(sys.argv) > 1 else "data/goldset"
LABELS = ["FIN_ORG", "ORG", "FIN_PRODUCT", "MONEY", "RATE", "DATE", "INDICATOR", "POLICY"]

def to_spans(text, cells, missing, who, sid):
    cands = []
    for tag, val in cells.items():
        for w in [x.strip() for x in str(val or "").split("|") if x.strip()]:
            i, found = text.find(w), False
            while i >= 0:
                cands.append((i, i + len(w), tag)); found = True; i = text.find(w, i + 1)
            if not found: missing.append({"who": who, "id": sid, "tag": tag, "text": w})
    keep = []
    for s, e, t in sorted(set(cands), key=lambda x: (-(x[1] - x[0]), x[0])):
        if all(e <= a or b <= s for a, b, _ in keep): keep.append((s, e, t))
    return sorted(keep)

def f1(a, b):
    a, b = set(a), set(b)
    if not a and not b: return 1.0
    tp = len(a & b); p = tp / len(b) if b else 0; r = tp / len(a) if a else 0
    return 2 * p * r / (p + r) if p + r else 0.0

def main():
    files = sorted(glob.glob(os.path.join(DIR, "gold_*_*.xlsx")))
    data = collections.defaultdict(dict)   # (set, id) -> {who: spans}
    texts, missing, unchecked = {}, [], collections.Counter()
    for fp in files:
        who = os.path.basename(fp)[:-5].split("_", 2)[2]
        wb = load_workbook(fp, read_only=True)
        for sheet in ("1_맞추기", "2_공통", "3_개인"):
            ws = wb[sheet]; rows = list(ws.iter_rows(values_only=True)); head = list(rows[0])
            for row in rows[1:]:
                if not row or row[1] is None: continue
                d = dict(zip(head, row)); sid, text = d["id"], d["문장"]; texts[sid] = text
                if str(d.get("확인") or "").strip().lower() not in ("o", "ㅇ", "v", "y", "1"): unchecked[who] += 1
                data[(sheet, sid)][who] = to_spans(text, {t: d.get(t) for t in LABELS}, missing, who, sid)
    report = {"files": [os.path.basename(f) for f in files], "unchecked_rows": dict(unchecked), "missing_text": missing[:200], "n_missing_text": len(missing)}
    # 일치율
    for sheet in ("1_맞추기", "2_공통"):
        items = [(sid, v) for (sh, sid), v in data.items() if sh == sheet]
        people = sorted({w for _, v in items for w in v})
        pair = {}
        for a, b in itertools.combinations(people, 2):
            pair[f"{a}~{b}"] = round(sum(f1(v.get(a, []), v.get(b, [])) for _, v in items) / max(len(items), 1), 4)
        by_tag = {}
        for t in LABELS:
            vals = [f1([x for x in v.get(a, []) if x[2] == t], [x for x in v.get(b, []) if x[2] == t])
                    for _, v in items for a, b in itertools.combinations(people, 2)]
            by_tag[t] = round(sum(vals) / max(len(vals), 1), 4)
        report[f"agreement_{sheet}"] = {"people": people, "mean_pairwise_f1": round(sum(pair.values()) / max(len(pair), 1), 4), "pairs": pair, "by_tag": by_tag}
    # 정답 합치기
    gold = []
    for (sheet, sid), v in sorted(data.items()):
        if sheet == "1_맞추기": continue
        if sheet == "2_공통":
            c = collections.Counter(x for spans in v.values() for x in spans)
            spans = sorted(x for x, k in c.items() if k >= 3)
        else:
            spans = next(iter(v.values()))
        gold.append({"id": sid, "category": "경제", "text": texts[sid], "spans": [list(x) for x in spans], "from": sheet})
    with open(os.path.join(DIR, "gold_final.jsonl"), "w", encoding="utf-8") as f:
        for r in gold: f.write(json.dumps(r, ensure_ascii=False) + "\n")
    report["gold_sentences"] = len(gold)
    report["gold_entities_by_tag"] = dict(collections.Counter(x[2] for r in gold for x in r["spans"]))
    json.dump(report, open(os.path.join(DIR, "import_report.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(json.dumps({k: v for k, v in report.items() if k != "missing_text"}, ensure_ascii=False, indent=1))

if __name__ == "__main__":
    main()
